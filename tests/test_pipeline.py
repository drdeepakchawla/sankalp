"""Unit tests on small synthetic data (no real records are kept in the repo)."""
import pandas as pd
import pytest

from sankalp import harmonise as H
from sankalp import link, outcomes, quality

AS_OF = pd.Timestamp("2026-10-13")


def test_phone10_keeps_valid_mobiles_only():
    s = pd.Series(["+91 98765 43210", "0", "1234567890", "9999999999", None, "09812345678"])
    out = H.phone10(s)
    assert out.iloc[0] == "9876543210"
    assert out.iloc[1:5].isna().all()
    assert out.iloc[5] == "9812345678"


def test_block_and_facility_names_are_harmonised():
    assert H.block(pd.Series(["Seagoan", "//Bhikangaon", "Not a usual resident of district", "MAHESHWAR", None])).tolist() == \
        ["Segaon", "Bhikangaon", "Out of district", "Maheshwar", "Unknown"]
    f = H.facility(pd.Series(["Jhirnya-CHC-Zirniya", "CHC- Zirniya", "Barwah-CH-Sanawad", "SDH- Sanawad", "Maheshwar-PHC-Chouli", "PHC-Choli"]))
    assert f.tolist() == ["CHC Zirniya", "CHC Zirniya", "SDH Sanawad", "SDH Sanawad", "PHC Choli", "PHC Choli"]


def _babies(rows):
    base = dict(period="2025-26", record_id="1", baby_no=1, born_status="Live born", outcome="Alive",
                dob=pd.Timestamp("2026-01-01"), dod=pd.NaT, fu_date=pd.Timestamp("2026-02-01"), bw_g=2800.0, ga_wk=39.0,
                preterm_reported=pd.NA, alive_d28_reported=pd.NA, fu_connected=True)
    return pd.DataFrame([{**base, **r} for r in rows])


def test_neonatal_death_window_and_denominators():
    b = _babies([
        dict(record_id="1"),                                                         # alive
        dict(record_id="2", outcome="Died", dod=pd.Timestamp("2026-01-01")),         # day 0 death
        dict(record_id="3", outcome="Died", dod=pd.Timestamp("2026-01-28")),         # day 27: neonatal
        dict(record_id="4", outcome="Died", dod=pd.Timestamp("2026-01-29")),         # day 28: post-neonatal
        dict(record_id="5", born_status="Stillborn", outcome="Stillbirth"),
        dict(record_id="6", outcome="Unknown - not reachable"),                      # due, unknown
        dict(record_id="7", outcome="Unknown - not followed up", dob=pd.Timestamp("2026-10-01")),  # not yet 28 days
    ])
    f = outcomes.finalise(b, AS_OF)
    assert f.loc[3, "outcome"] == "Died after day 27"
    assert f.loc[6, "outcome"] == "Pending (under 28 days)"
    r = outcomes.nmr(f)
    assert r["deaths"] == 2 and r["early"] == 1 and r["late"] == 1
    assert r["known"] == 4            # alive + 2 neonatal deaths + 1 post-neonatal death
    assert r["live_births"] == 6 and r["stillbirths"] == 1 and r["births"] == 7
    assert r["nmr"] == pytest.approx(500.0)


def test_linkage_prefers_phone_then_names_and_rejects_far_dates():
    babies = pd.DataFrame({
        "baby_uid": ["a", "b", "c"], "dob": pd.to_datetime(["2025-03-01", "2025-03-01", "2025-06-01"]),
        "phone1": ["9876543210", None, "9000000001"], "phone2": [None, None, None],
        "sex": ["Male", "Female", "Male"], "bw_g": [2500.0, 1800.0, 3000.0],
        "mother_name": ["Sunita", "Rekha Bai", "Kavita"], "father_name": ["Ramesh", "Mohan Lal", "Suresh"]})
    sncu = pd.DataFrame({
        "dob": pd.to_datetime(["2025-03-02", "2025-03-01", "2025-07-15"]), "phone": ["9876543210", None, "9000000001"],
        "sex": ["Male", "Female", "Male"], "bw_g": [2450.0, 1750.0, 3000.0],
        "mother_name": ["Sunita", "Rekha", "Kavita"], "father_name": ["Ramesh Patel", "Mohanlal", "Suresh"]})
    out = link.link(babies, sncu)
    assert out.loc[0, "baby_uid"] == "a" and out.loc[0, "link_tier"] == 1
    assert out.loc[1, "baby_uid"] == "b" and out.loc[1, "link_tier"] == 2
    assert pd.isna(out.loc[2, "baby_uid"])   # same phone but DOB 44 days apart


def test_duplicate_needs_more_than_a_shared_phone():
    b = pd.DataFrame({
        "period": ["2025-26"] * 4, "record_id": ["1", "2", "3", "4"], "baby_no": [1] * 4,
        "dob": pd.to_datetime(["2025-05-05"] * 4), "phone1": ["9876543210", "9876543210", "9123456789", "9123456789"],
        "mother_name": ["Pooja", "Pooja Bai", "Anita", "Savita"], "father_name": ["Raju", "Raju", "Arun", "Vijay"]})
    g = quality.duplicate_groups(b)
    assert g.iloc[0] == g.iloc[1] and pd.notna(g.iloc[0])   # same phone, same mother ("Bai" is ignored)
    assert g.iloc[2:].isna().all()                          # same phone (e.g. an ASHA's), different mothers


def test_error_workbook_has_a_sheet_per_check(tmp_path):
    from openpyxl import load_workbook
    from sankalp import error_listing
    babies = pd.DataFrame({
        "baby_uid": ["2025-26-1-1", "2025-26-2-1", "2025-26-3-1"], "period": ["2025-26"] * 3, "record_id": ["1", "2", "3"],
        "baby_no": [1, 1, 1], "mother_name": ["Asha", "Asha", "Meena"], "father_name": ["Ravi", "Ravi", "Raju"],
        "phone1": ["9876543210", "9876543210", None], "phone2": [None] * 3, "village": ["A", "A", "B"], "res_block": ["Barwah"] * 3,
        "dob": pd.to_datetime(["2025-05-05", "2025-05-05", "2025-06-01"]), "birth_place_type": ["Public facility"] * 3,
        "birth_facility": ["DH Khargone"] * 3, "sex": ["Male", "Male", "Female"], "bw_g": [2500.0, 2500.0, 150.0], "ga_wk": [39.0, 39.0, 38.0],
        "outcome": ["Alive", "Alive", "Alive"], "dod": pd.NaT, "fu_date": pd.to_datetime(["2025-06-10"] * 3), "fu_days": [36, 36, 9],
        "reg_collector_name": ["DC A"] * 3, "fu_collector_name": ["DC B"] * 3, "birth_info_source": ["Facility data"] * 3,
        "dup_group": ["N1", "N1", None], "status_detail": ["Day-28 call: Healthy"] * 3, "lmp": pd.NaT, "age_death_days": float("nan"),
        "flags": [["DUPLICATE_SUSPECT"], ["DUPLICATE_SUSPECT"], ["BW_IMPLAUSIBLE"]]})
    sncu = pd.DataFrame({
        "admission_id": ["IND/KHR/SNCU1/2025/1"], "adm_date": pd.to_datetime(["2025-06-02"]), "dob": pd.to_datetime(["2025-06-01"]),
        "mother_name": ["Meena"], "father_name": ["Raju"], "phone": [None], "village": ["B"], "res_block": ["Barwah"], "sex": ["Female"],
        "bw_g": [1500.0], "ga_wk": [32.0], "admission_type": ["Inborn"], "place_of_delivery": ["DH Khargone"], "outcome": ["EXPIRED"],
        "discharge_date": pd.to_datetime(["2025-06-04"]), "cause_or_reason": ["Sepsis"], "cause_of_death": ["Sepsis"],
        "baby_uid": ["2025-26-3-1"], "link_tier": [2.0], "maturity": ["Preterm"], "age_at_adm_days": [1.0], "flags": [["S_DIED_COHORT_ALIVE"]]})
    out = tmp_path / "errors.xlsx"
    res = error_listing.write(babies, sncu, out, "2026-10-13")
    wb = load_workbook(out)
    assert res["issues"] == 4
    assert [n for n in wb.sheetnames if n[:2].isdigit()] == ["01 Possible duplicates", "02 Birth weight implausible", "03 SNCU died, cohort alive"]
    dup = wb["01 Possible duplicates"]
    assert dup.max_row == 3 and dup.freeze_panes == "E2"
    assert "Same baby as: 2025-26 Record 2 baby 1" in [c.value for c in dup[2]]
