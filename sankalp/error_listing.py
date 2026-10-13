"""Excel line list of every data-quality issue, for correcting records in REDCap.

One sheet per check (identifiers included), plus a summary, staff counts, a long list of
every issue and a list by record. Identifiable: keep it inside the Sankalp team.
"""
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList
from rapidfuzz import fuzz, process

from . import quality as Q
from .link import _norm, name_sim

FONT = "Arial"
SEV_RANK = {"critical": 0, "major": 1, "minor": 2}
SEV_FILL = {"critical": "F6D5D3", "major": "FBE8BE", "minor": "E6E9E7"}
HEAD_FILL, HEAD_FONT = "1E5D63", "FFFFFF"
INPUT_FILL = "FFF2CC"  # columns the team fills in
FILL_IN = [("Action taken", 34), ("Corrected on", 14), ("Corrected by", 18), ("Remarks", 34)]
DATE_FMT = "dd-mmm-yyyy"
PERIOD_PROJECT = {"2024": "Sankalp Khargone REDCap project", "2025-26": "NHRP Sankalp REDCap project"}

BABY_COLS = [  # (header, column, width)
    ("Period", "period", 9), ("Record ID", "record_id", 12), ("Baby no.", "baby_no", 7),
    ("Mother name", "mother_name", 18), ("Father name", "father_name", 18), ("Phone", "phone1", 12), ("Phone 2", "phone2", 12),
    ("Village", "village", 16), ("Block", "res_block", 13), ("Date of birth", "dob", 12),
    ("Place of birth", "birth_place_type", 14), ("Facility of birth", "birth_facility", 18), ("Sex", "sex", 8),
    ("Birth weight (g)", "bw_g", 10), ("Gestation (wk)", "ga_wk", 10), ("Day-28 outcome", "outcome", 20),
    ("Date of death", "dod", 12), ("Outcome date", "fu_date", 12), ("Registered by", "reg_collector_name", 22),
    ("Outcome collected by", "fu_collector_name", 20), ("Data source", "birth_info_source", 18),
]
SNCU_COLS = [
    ("Admission ID", "admission_id", 24), ("Admission date", "adm_date", 12), ("Date of birth", "dob", 12),
    ("Mother name", "mother_name", 18), ("Father name", "father_name", 18), ("Phone", "phone", 12), ("Village", "village", 16),
    ("Block", "res_block", 13), ("Sex", "sex", 8), ("Weight (g)", "bw_g", 10), ("Gestation (wk)", "ga_wk", 10),
    ("Type of admission", "admission_type", 20), ("Place of delivery", "place_of_delivery", 18), ("SNCU outcome", "outcome", 13),
    ("Discharge date", "discharge_date", 12), ("Cause of death / referral reason", "cause_or_reason", 26),
    ("Cohort record", "cohort_ref", 24), ("Cohort outcome", "cohort_outcome", 20), ("Cohort outcome date", "cohort_fu_date", 12),
]
ALL_COLS = [("Dataset", 10), ("Period", 9), ("Record / admission ID", 22), ("Baby no.", 7), ("Severity", 9), ("Check code", 26),
            ("Check", 34), ("Block", 13), ("Mother name", 18), ("Phone", 12), ("Registered by", 22), ("Outcome collected by", 20),
            ("Issue detail", 60), ("Sheet", 30)]
REC_COLS = [("Dataset", 10), ("Period", 9), ("Record / admission ID", 22), ("Baby no.", 7), ("Mother name", 18), ("Father name", 18),
            ("Phone", 12), ("Village", 16), ("Block", 13), ("Date of birth", 12), ("Registered by", 22), ("Outcome collected by", 20),
            ("Number of issues", 9), ("Most severe", 10), ("Issues", 70)]


# ---------- helpers ----------
def _d(v):
    return "" if v is None or pd.isna(v) else pd.Timestamp(v).strftime("%d-%m-%Y")


def _n(v):
    return "" if v is None or pd.isna(v) else f"{v:,.0f}"


def _clean(v):
    if v is None or v is pd.NaT:
        return None
    if isinstance(v, (list, tuple)):
        return "; ".join(map(str, v))
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, (np.floating, float)):
        return int(v) if float(v).is_integer() else round(float(v), 1)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.bool_, bool)):
        return "Yes" if v else "No"
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def _lists(s):
    return s.map(lambda v: [x for x in v.split("; ") if x] if isinstance(v, str) else list(v) if isinstance(v, (list, tuple, np.ndarray)) else [])


class Book:
    """Thin wrapper over a write-only openpyxl workbook with Arial everywhere."""

    def __init__(self):
        self.wb = Workbook(write_only=True)
        self.wb._fonts = IndexedList([Font(name=FONT, sz=10, family=2)])  # default cell font
        self.f = Font(name=FONT, sz=10)
        self.fb = Font(name=FONT, sz=10, bold=True)
        self.fh = Font(name=FONT, sz=10, bold=True, color=HEAD_FONT)
        self.ft = Font(name=FONT, sz=14, bold=True)
        self.fl = Font(name=FONT, sz=10, color="0563C1", underline="single")
        self.fin = Font(name=FONT, sz=10, color="0000FF")
        self.wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(self, title, widths, freeze=None):
        ws = self.wb.create_sheet(title)
        if freeze:
            ws.freeze_panes = freeze  # write-only sheets keep this only if set before the first row
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    def cell(self, ws, v, font=None, fill=None, fmt=None, wrap=False):
        c = WriteOnlyCell(ws, value=_clean(v))
        c.font = font or self.f
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
        if fmt:
            c.number_format = fmt
        if wrap:
            c.alignment = self.wrap
        return c

    def header(self, ws, names, inputs=()):
        ws.append([self.cell(ws, n, self.fb if n in inputs else self.fh, INPUT_FILL if n in inputs else HEAD_FILL, wrap=True) for n in names])

    def table(self, ws, df, cols, date_cols=(), wrap_cols=(), inputs=True):
        """Write a header + rows. cols: list of (header, df column, width)."""
        names = [h for h, _, _ in cols] + ([h for h, _ in FILL_IN] if inputs else [])
        self.header(ws, names, inputs={h for h, _ in FILL_IN} if inputs else ())
        keys = [c for _, c, _ in cols]
        vals = df[keys].to_numpy(dtype=object) if len(df) else np.empty((0, len(keys)), dtype=object)
        for row in vals:
            out = []
            for k, v in zip(keys, row):
                if k in date_cols:
                    out.append(self.cell(ws, v, fmt=DATE_FMT))
                elif k in wrap_cols:
                    out.append(self.cell(ws, v, wrap=True))
                else:
                    out.append(_clean(v))
            ws.append(out)
        last = get_column_letter(len(names))
        ws.auto_filter.ref = f"A1:{last}{max(2, len(df) + 1)}"


# ---------- per-check detail ----------
def _baby_detail(code, r, as_of, sncu_by_baby):
    trail = r.get("status_detail") if isinstance(r.get("status_detail"), str) else ""
    if code == "OUTCOME_UNKNOWN":
        age = (as_of - r["dob"]).days if pd.notna(r["dob"]) else None
        return f"{r['outcome']}; baby is {age} days old. Forms: {trail}" if age is not None else f"{r['outcome']}. Forms: {trail}"
    if code in ("STATUS_CONFLICT", "STILLBIRTH_ALIVE_LATER", "DEATH_DATE_CONFLICT", "DEATH_AFTER_CALL", "BIRTH_STATUS_MISSING"):
        return f"Forms: {trail}"
    if code == "DEATH_DATE_MISSING":
        return f"Died, no date of death. Forms: {trail}"
    if code == "DEATH_BEFORE_BIRTH":
        return f"Date of death {_d(r['dod'])} is before date of birth {_d(r['dob'])}"
    if code == "FU_BEFORE_BIRTH":
        return f"Outcome date {_d(r['fu_date'])} is before date of birth {_d(r['dob'])}"
    if code in ("SNCU_DEATH_NOT_RECORDED", "SNCU_STILLBIRTH"):
        adm = sncu_by_baby.get(r["baby_uid"], [])
        txt = "; ".join(f"SNCU {a['admission_id']} admitted {_d(a['adm_date'])}, {a['outcome']} {_d(a['discharge_date'])}"
                        + (f" ({a['cause_of_death']})" if isinstance(a.get("cause_of_death"), str) else "") for a in adm)
        lead = "Recorded stillborn, but admitted to SNCU" if code == "SNCU_STILLBIRTH" else f"Cohort outcome '{r['outcome']}', but"
        return f"{lead}: {txt}. Forms: {trail}"
    if code == "DOB_MISSING":
        return "Date of birth not recorded"
    if code == "DOB_FUTURE":
        return f"Date of birth {_d(r['dob'])} is after the export date {_d(as_of)}"
    if code == "DIED_AFTER_D27_MARKED_DEAD":
        return f"'Alive on day 28' = No, but death on day {_n(r['age_death_days'])} ({_d(r['dod'])})"
    if code == "FU_EARLY":
        return f"Reported alive at a call on day {_n(r['fu_days'])} ({_d(r['fu_date'])}); day-28 status not confirmed"
    if code == "FU_LATE":
        return f"Outcome collected on day {_n(r['fu_days'])} ({_d(r['fu_date'])})"
    if code == "BW_MISSING":
        return "Birth weight not recorded"
    if code == "BW_IMPLAUSIBLE":
        return f"Birth weight recorded as {_n(r['bw_g'])} g"
    if code == "GA_IMPLAUSIBLE":
        return f"Gestation {_n(r['ga_wk'])} weeks" + (f" from LMP {_d(r['lmp'])}" if pd.notna(r.get("lmp")) else "")
    if code == "GA_MISSING":
        return "Gestation not recorded (LMP not known)"
    if code == "BW_GA_INCONSISTENT":
        return f"{_n(r['bw_g'])} g at {_n(r['ga_wk'])} weeks" + (f" (LMP {_d(r['lmp'])})" if pd.notna(r.get("lmp")) else "")
    if code == "SEX_MISSING":
        return "Sex not recorded"
    if code == "SEX_CONFLICT":
        return f"Notification: {r.get('notif_sex')}; registration: {r['sex']}"
    if code == "BW_CONFLICT":
        return f"Registration {_n(r['bw_g'])} g; day-28 call {_n(r.get('fu_bw'))} g"
    if code == "BLOCK_MISSING":
        return "Block of residence not recorded"
    if code == "DEATH_PLACE_MISSING":
        return "Place of death not recorded"
    if code == "DUPLICATE_SUSPECT":
        return "Same baby as:"
    return Q.BABY_CHECKS.get(code, ("", code, ""))[2]


def _sncu_detail(code, r, last_adm):
    if code == "S_ADMISSION_BEFORE_BIRTH":
        return f"Admitted {_d(r['adm_date'])}, before date of birth {_d(r['dob'])}"
    if code == "S_DIED_COHORT_ALIVE":
        trail = r.get("cohort_trail") if isinstance(r.get("cohort_trail"), str) else ""
        return f"SNCU death on {_d(r['discharge_date'])}; cohort says '{r['cohort_outcome']}'. Cohort forms: {trail}"
    if code == "S_OUTCOME_MISSING":
        recent = pd.notna(r["adm_date"]) and (last_adm - r["adm_date"]).days <= 30
        return "No outcome recorded" + (" (admitted in the last 30 days of the register: may still be in the SNCU)" if recent else "")
    if code == "S_DEATH_NOT_IN_DEATH_SHEET":
        return f"Outcome EXPIRED on {_d(r['discharge_date'])}, but not in the deaths sheet"
    if code == "S_DISCHARGE_BEFORE_ADMISSION":
        return f"Discharged {_d(r['discharge_date'])}, admitted {_d(r['adm_date'])}"
    if code == "S_WEIGHT_IMPLAUSIBLE":
        return f"Weight recorded as {_n(r['bw_g'])} g"
    if code == "S_NOT_IN_COHORT":
        cand = r.get("candidate")
        return cand if isinstance(cand, str) else "No similar cohort record (date of birth within 10 days, same phone or similar mother's name)"
    if code == "S_MARKED_DUPLICATE":
        return "Marked 'Duplicate' in the register"
    if code == "S_AGE_MISMATCH":
        calc = (r["adm_date"] - r["dob"]).days if pd.notna(r["adm_date"]) and pd.notna(r["dob"]) else None
        return f"Recorded age {_n(r['age_at_adm_days'])} days; admission date minus date of birth = {calc} days"
    if code == "S_GA_MATURITY_MISMATCH":
        return f"Maturity '{r['maturity']}' but gestation {_n(r['ga_wk'])} weeks"
    return Q.SNCU_CHECKS.get(code, ("", code, ""))[2]


def _same_phone(ph, *others):
    return isinstance(ph, str) and any(isinstance(o, str) and o == ph for o in others)


def _candidates(unmatched, babies):
    """Closest cohort baby for SNCU admissions that did not link, for a person to confirm."""
    b = babies[babies["dob"].notna()].sort_values("dob").reset_index(drop=True)
    days = b["dob"].to_numpy().astype("datetime64[D]").astype(np.int64)
    m, f = _norm(b["mother_name"]).to_numpy(), _norm(b["father_name"]).to_numpy()
    by_phone = {}
    for col in ("phone1", "phone2"):
        for pos, ph in enumerate(b[col]):
            if isinstance(ph, str):
                by_phone.setdefault(ph, []).append(pos)
    out = {}
    for i, a in unmatched.iterrows():
        if pd.isna(a["dob"]):
            continue
        d = np.datetime64(a["dob"], "D").astype(np.int64)
        am, af = _norm(pd.Series([a["mother_name"]])).iat[0], _norm(pd.Series([a["father_name"]])).iat[0]
        lo, hi = np.searchsorted(days, d - 10, "left"), np.searchsorted(days, d + 10, "right")
        cands = set()
        if am:
            cands |= {lo + j for _, _, j in process.extract(am, m[lo:hi], scorer=fuzz.token_set_ratio, score_cutoff=60, limit=15)}
        if isinstance(a["phone"], str):
            cands |= {p for p in by_phone.get(a["phone"], []) if abs(days[p] - d) <= 60}
        best, best_s = None, -1e9
        for p in cands:
            ph = _same_phone(a["phone"], b.at[p, "phone1"], b.at[p, "phone2"])
            s = name_sim(am, m[p]) + name_sim(af, f[p]) + 60 * ph - 3 * abs(int(days[p] - d))
            if s > best_s:
                best, best_s = p, s
        if best is not None:
            r = b.loc[best]
            ph = _same_phone(a["phone"], r["phone1"], r["phone2"])
            out[i] = (f"Possible match: {r['period']} Record {r['record_id']} baby {r['baby_no']} - mother '{r['mother_name']}' "
                      f"({name_sim(am, m[best]):.0f}%), father '{r['father_name']}' ({name_sim(af, f[best]):.0f}%), "
                      f"DOB {_d(r['dob'])} ({int(days[best] - d):+d} days), {'same phone' if ph else 'different phone'}")
    return out


SHEET_NAMES = {
    "OUTCOME_UNKNOWN": "Outcome unknown", "FU_BEFORE_BIRTH": "Follow-up before birth", "STILLBIRTH_ALIVE_LATER": "Stillbirth then alive",
    "SNCU_DEATH_NOT_RECORDED": "SNCU death not in cohort", "STATUS_CONFLICT": "Alive after death", "SNCU_STILLBIRTH": "Stillbirth in SNCU",
    "DEATH_BEFORE_BIRTH": "Death before birth", "DOB_MISSING": "Birth date missing", "DOB_FUTURE": "Birth date in future",
    "DUPLICATE_SUSPECT": "Possible duplicates", "FU_EARLY": "Follow-up before day 28", "GA_IMPLAUSIBLE": "Gestation implausible",
    "BW_MISSING": "Birth weight missing", "BIRTH_STATUS_MISSING": "Birth status missing", "BW_IMPLAUSIBLE": "Birth weight implausible",
    "DEATH_AFTER_CALL": "Death after call", "DIED_AFTER_D27_MARKED_DEAD": "Death after day 27", "DEATH_DATE_CONFLICT": "Death dates disagree",
    "DEATH_DATE_MISSING": "Death date missing", "GA_MISSING": "Gestation missing", "FU_LATE": "Follow-up after day 90",
    "BLOCK_MISSING": "Block missing", "SEX_MISSING": "Sex missing", "BW_CONFLICT": "Weight differs", "BW_GA_INCONSISTENT": "Weight vs gestation",
    "SEX_CONFLICT": "Sex differs", "DEATH_PLACE_MISSING": "Death place missing",
    "S_ADMISSION_BEFORE_BIRTH": "SNCU admitted before birth", "S_DIED_COHORT_ALIVE": "SNCU died, cohort alive", "S_OUTCOME_MISSING": "SNCU outcome missing",
    "S_DEATH_NOT_IN_DEATH_SHEET": "SNCU death not in sheet", "S_DISCHARGE_BEFORE_ADMISSION": "SNCU discharge before adm",
    "S_WEIGHT_IMPLAUSIBLE": "SNCU weight implausible", "S_NOT_IN_COHORT": "SNCU not in cohort", "S_MARKED_DUPLICATE": "SNCU marked duplicate",
    "S_AGE_MISMATCH": "SNCU age mismatch", "S_GA_MATURITY_MISMATCH": "SNCU GA vs maturity",
}


def _sheet_name(i, label, used):
    bad = '[]:*?/\\'
    base = "".join(ch for ch in label if ch not in bad).replace("'", "")
    name = f"{i:02d} {base}"[:31].rstrip()
    while name in used:
        name = name[:-1]
    used.add(name)
    return name


# ---------- workbook ----------
def write(babies: pd.DataFrame, sncu: pd.DataFrame, path, as_of) -> dict:
    as_of = pd.Timestamp(as_of)
    b = babies.copy()
    s = sncu.copy()
    b["flag_list"] = _lists(b["flags"])
    s["flag_list"] = _lists(s["flags"])
    for c in ("reg_collector_name", "fu_collector_name", "status_detail", "notif_sex", "fu_bw", "lmp"):
        if c not in b:
            b[c] = pd.NA
    s["cause_or_reason"] = s.get("cause_or_reason", s.get("cause_of_death"))

    bi = b.set_index("baby_uid")
    s["cohort_ref"] = s["baby_uid"].map(lambda u: f"{bi.at[u, 'period']} · Record {bi.at[u, 'record_id']} · baby {bi.at[u, 'baby_no']}" if isinstance(u, str) else None)
    s["cohort_outcome"] = s["baby_uid"].map(bi["outcome"])
    s["cohort_fu_date"] = s["baby_uid"].map(bi["fu_date"])
    s["cohort_trail"] = s["baby_uid"].map(bi["status_detail"])
    unmatched = s[s["flag_list"].map(lambda l: "S_NOT_IN_COHORT" in l)]
    s["candidate"] = pd.Series(_candidates(unmatched, b), dtype=object).reindex(s.index)
    sncu_by_baby = {u: g.to_dict("records") for u, g in s.dropna(subset=["baby_uid"]).sort_values("adm_date").groupby("baby_uid")}
    last_adm = s["adm_date"].max()

    # Duplicate groups: list the other records in each group.
    others = {}
    for g, grp in b.dropna(subset=["dup_group"]).groupby("dup_group"):
        refs = [(i, f"{r.period} Record {r.record_id} baby {r.baby_no}") for i, r in grp.iterrows()]
        for i, _ in refs:
            others[i] = ", ".join(t for j, t in refs if j != i)

    bk = Book()
    checks = [("Cohort", code, *Q.BABY_CHECKS[code]) for code in Q.BABY_CHECKS] + [("SNCU", code, *Q.SNCU_CHECKS[code]) for code in Q.SNCU_CHECKS]
    counts = {code: int((b if ds == "Cohort" else s)["flag_list"].map(lambda l, c=code: c in l).sum()) for ds, code, *_ in checks}
    checks = [c for c in checks if counts[c[1]]]
    checks.sort(key=lambda c: (c[0] != "Cohort", SEV_RANK[c[2]], -counts[c[1]]))
    used = {"Read me", "Summary", "By staff", "All errors", "By record"}
    sheet_of = {c[1]: _sheet_name(n + 1, SHEET_NAMES.get(c[1], c[3]), used) for n, c in enumerate(checks)}

    long_rows = []
    rec = {}

    def other_issues(flags, code, table):
        return "; ".join(table[c][1] for c in flags if c != code and c in table)

    _readme(bk, as_of, checks, counts, sheet_of)
    _summary(bk, checks, sheet_of)
    staff_ws = bk.sheet("By staff", [9, 34, 12, 12, 10, 10, 10, 12], freeze="A3")

    for ds, code, sev, label, help_ in checks:
        if ds == "Cohort":
            rows = b[b["flag_list"].map(lambda l, c=code: c in l)].copy()
            rows["detail"] = [(_baby_detail(code, r, as_of, sncu_by_baby) + (f" {others.get(i, '')}" if code == "DUPLICATE_SUSPECT" else ""))
                              for i, r in rows.iterrows()]
            rows["other"] = rows["flag_list"].map(lambda l, c=code: other_issues(l, c, Q.BABY_CHECKS))
            cols = BABY_COLS + [("Issue detail", "detail", 60), ("Other issues", "other", 34)]
            if code == "DUPLICATE_SUSPECT":
                rows = rows.sort_values(["dup_group", "period", "record_id"])
                cols = [("Duplicate group", "dup_group", 11)] + cols
            else:
                rows = rows.sort_values(["period", "res_block", "dob"])
            ws = bk.sheet(sheet_of[code], [w for _, _, w in cols] + [w for _, w in FILL_IN], freeze="E2" if code == "DUPLICATE_SUSPECT" else "D2")
            bk.table(ws, rows, cols, date_cols={"dob", "dod", "fu_date"}, wrap_cols={"detail", "other"})
            for i, r in rows.iterrows():
                long_rows.append(["Cohort", r["period"], r["record_id"], r["baby_no"], sev, code, label, r["res_block"], r["mother_name"],
                                  r["phone1"], r["reg_collector_name"], r["fu_collector_name"], r["detail"], sheet_of[code]])
                e = rec.setdefault(("Cohort", i), {"r": r, "codes": []})
                e["codes"].append(code)
        else:
            rows = s[s["flag_list"].map(lambda l, c=code: c in l)].copy()
            rows["detail"] = [_sncu_detail(code, r, last_adm) for _, r in rows.iterrows()]
            rows["other"] = rows["flag_list"].map(lambda l, c=code: other_issues(l, c, Q.SNCU_CHECKS))
            rows = rows.sort_values(["adm_date"])
            cols = SNCU_COLS + [("Issue detail", "detail", 60), ("Other issues", "other", 34)]
            ws = bk.sheet(sheet_of[code], [w for _, _, w in cols] + [w for _, w in FILL_IN], freeze="B2")
            bk.table(ws, rows, cols, date_cols={"adm_date", "dob", "discharge_date", "cohort_fu_date"}, wrap_cols={"detail", "other"})
            for i, r in rows.iterrows():
                yr = r["adm_date"].strftime("%Y") if pd.notna(r["adm_date"]) else None
                long_rows.append(["SNCU", yr, r["admission_id"], None, sev, code, label, r["res_block"], r["mother_name"], r["phone"],
                                  None, None, r["detail"], sheet_of[code]])
                e = rec.setdefault(("SNCU", i), {"r": r, "codes": []})
                e["codes"].append(code)

    # All errors (long)
    ws = bk.sheet("All errors", [w for _, w in ALL_COLS], freeze="D2")
    bk.header(ws, [h for h, _ in ALL_COLS])
    for row in long_rows:
        ws.append([bk.cell(ws, v, wrap=True) if j == 12 else _clean(v) for j, v in enumerate(row)])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(ALL_COLS))}{len(long_rows) + 1}"

    # By record
    ws = bk.sheet("By record", [w for _, w in REC_COLS] + [w for _, w in FILL_IN], freeze="D2")
    bk.header(ws, [h for h, _ in REC_COLS] + [h for h, _ in FILL_IN], inputs={h for h, _ in FILL_IN})
    labels = {**{k: v for k, v in Q.BABY_CHECKS.items()}, **Q.SNCU_CHECKS}
    order = sorted(rec.items(), key=lambda kv: (kv[0][0] != "Cohort", min(SEV_RANK[labels[c][0]] for c in kv[1]["codes"]),
                                                str(kv[1]["r"].get("period", "")), str(kv[1]["r"].get("res_block", ""))))
    for (ds, _), e in order:
        r, codes = e["r"], sorted(e["codes"], key=lambda c: SEV_RANK[labels[c][0]])
        worst = labels[codes[0]][0]
        common = [r["mother_name"], r["father_name"], r["phone1"] if ds == "Cohort" else r["phone"], r["village"], r["res_block"]]
        if ds == "Cohort":
            head = ["Cohort", r["period"], r["record_id"], r["baby_no"]] + common + [bk.cell(ws, r["dob"], fmt=DATE_FMT), r["reg_collector_name"], r["fu_collector_name"]]
        else:
            yr = r["adm_date"].strftime("%Y") if pd.notna(r["adm_date"]) else None
            head = ["SNCU", yr, r["admission_id"], None] + common + [bk.cell(ws, r["dob"], fmt=DATE_FMT), None, None]
        head = [v if isinstance(v, Cell) else _clean(v) for v in head]
        ws.append(head + [len(codes), bk.cell(ws, worst, fill=SEV_FILL[worst]), bk.cell(ws, "; ".join(labels[c][1] for c in codes), wrap=True)])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(REC_COLS) + len(FILL_IN))}{len(rec) + 1}"

    _staff(bk, staff_ws, b)
    bk.wb.save(path)
    return {"checks": len(checks), "issues": len(long_rows), "records": len(rec), "sheets": len(checks) + 5}


def _readme(bk, as_of, checks, counts, sheet_of):
    ws = bk.sheet("Read me", [26, 100])
    c = lambda v, **k: bk.cell(ws, v, **k)
    ws.append([c("Sankalp data error line list", font=bk.ft)])
    ws.append([c(f"Data as of {as_of.strftime('%d %b %Y')} · generated {datetime.now().strftime('%d %b %Y %H:%M')}")])
    ws.append([])
    ws.append([c("Confidential", font=bk.fb, fill=SEV_FILL["critical"]),
               c("Contains names, phone numbers and villages of mothers and babies, and names of staff. Keep it within the Sankalp team. "
                 "Do not forward or upload it anywhere else.", fill=SEV_FILL["critical"], wrap=True)])
    ws.append([])
    rows = [
        ("What this is", "Every record that fails a data-quality check, from the 2024 Sankalp Khargone cohort, the 2025-26 NHRP cohort and the DH Khargone SNCU register."),
        ("Summary", "Each check with its severity and the number of records per period, with a link to its sheet."),
        ("Numbered sheets", "One sheet per check (01, 02 ...), most serious first. Each row is one baby (or SNCU admission) with its identifiers, "
                            "the exact problem in 'Issue detail', and any other checks the same record fails."),
        ("All errors", "One row per record per issue, for filtering and pivot tables."),
        ("By record", "One row per record with all its issues together, so a record can be fixed in one call or visit."),
        ("By staff", "Records with issues by the staff member who registered the birth and who collected the outcome."),
        ("Finding the record", "Period 2024 = Sankalp Khargone REDCap project; 2025-26 = NHRP Sankalp REDCap project. Search by Record ID. "
                               "SNCU rows: search the SNCU register by Admission ID."),
        ("Severity", "critical = changes the count of neonatal deaths or the NMR; major = changes a key measure (weight, gestation, dates, duplicates); "
                     "minor = a missing or inconsistent detail."),
        ("Day-28 outcome", "As worked out by the pipeline from all forms: Alive, Died (0-27 days), Died after day 27, Stillbirth, "
                           "Pending (under 28 days old) or Unknown (not reachable / not followed up / not known)."),
        ("Issue detail · 'Forms:'", "What each form reported, in order, e.g. 'Registration: Alive on 02-03-2025; Day-28 call: Dead on 05-04-2025, died 20-03-2025'."),
    ]
    for k, v in rows:
        ws.append([c(k, font=bk.fb, wrap=True), c(v, wrap=True)])
    ws.append([])
    ws.append([c("Columns to fill in", font=bk.fb, fill=INPUT_FILL),
               c("The yellow columns at the right of every check sheet and of 'By record': Action taken, Corrected on, Corrected by, Remarks. "
                 "Leave the other columns as they are; the next export will show whether the issue is gone.", fill=INPUT_FILL, wrap=True)])
    ws.append([c("Example of a filled row", font=bk.fb, wrap=True),
               c("Action taken: Corrected date of death in REDCap  |  Corrected on: 15-Oct-2026  |  Corrected by: Data manager  |  Remarks: Confirmed with ASHA by phone", wrap=True)])
    ws.append([])
    ws.append([c("Checks in this file", font=bk.fb), c(f"{len(checks)} checks with at least one record; {sum(counts[x[1]] for x in checks):,} issues in all.")])


def _summary(bk, checks, sheet_of):
    widths = [9, 10, 26, 34, 60, 30, 10, 10, 10, 10]
    ws = bk.sheet("Summary", widths, freeze="A4")
    c = lambda v, **k: bk.cell(ws, v, **k)
    ws.append([c("Issues by check", font=bk.ft)])
    ws.append([c("Counts are formulas over the 'All errors' sheet.", font=Font(name=FONT, sz=9, italic=True))])
    head = ["Dataset", "Severity", "Check code", "Check", "What it means", "Sheet", "2024", "2025-26", "SNCU", "Total"]
    bk.header(ws, head)
    first = 4
    for n, (ds, code, sev, label, help_) in enumerate(checks):
        r = first + n
        crit = lambda p: f"=COUNTIFS('All errors'!$F:$F,$C{r},'All errors'!$A:$A,\"Cohort\",'All errors'!$B:$B,\"{p}\")"
        ws.append([ds, c(sev, fill=SEV_FILL[sev]), c(code, font=Font(name="Courier New", sz=9)), label, c(help_, wrap=True),
                   c(f"=HYPERLINK(\"#'{sheet_of[code]}'!A1\",\"{sheet_of[code]}\")", font=bk.fl),
                   crit("2024"), crit("2025-26"), f"=COUNTIFS('All errors'!$F:$F,$C{r},'All errors'!$A:$A,\"SNCU\")", f"=SUM(G{r}:I{r})"])
    last = first + len(checks) - 1
    ws.append([c("Total issues", font=bk.fb), None, None, None, None, None] +
              [c(f"=SUM({col}{first}:{col}{last})", font=bk.fb) for col in "GHIJ"])
    ws.append([])
    ws.append([c("Records with issues", font=bk.fb)])
    for lab, ds, sev in [("Cohort babies with a critical issue", "Cohort", "critical"), ("Cohort babies with a major issue (none critical)", "Cohort", "major"),
                         ("Cohort babies with only minor issues", "Cohort", "minor"), ("SNCU admissions with any issue", "SNCU", None)]:
        f = (f"=COUNTIFS('By record'!$A:$A,\"{ds}\",'By record'!$N:$N,\"{sev}\")" if sev else f"=COUNTIFS('By record'!$A:$A,\"{ds}\")")
        ws.append([c(lab), None, None, None, None, None, None, None, None, f])


def _staff(bk, ws, b):
    """Per staff member: babies handled (input from the exports) and records with issues (formulas over 'By record')."""
    c = lambda v, **k: bk.cell(ws, v, **k)
    ws.append([c("Records with issues by staff member", font=bk.ft)])
    ws.append([c("'Babies' (blue) is the number of babies each person handled, counted from the exports. "
                 "The other columns are formulas over the 'By record' sheet.", font=Font(name=FONT, sz=9, italic=True))])
    row = 2
    for title, col_name, rec_col in [("Registered the birth", "reg_collector_name", "K"), ("Collected the day-28 outcome", "fu_collector_name", "L")]:
        ws.append([]); ws.append([c(title, font=bk.fb)])
        bk.header(ws, ["Period", "Staff", "Babies", "With any issue", "Critical", "Major", "Minor", "Share with any issue"])
        row += 3
        counts = (b.dropna(subset=[col_name]).groupby(["period", col_name]).size().reset_index(name="n")
                  .sort_values(["period", "n"], ascending=[True, False]))
        for _, r in counts.iterrows():
            row += 1
            base = f"'By record'!$A:$A,\"Cohort\",'By record'!$B:$B,$A{row},'By record'!${rec_col}:${rec_col},$B{row}"
            ws.append([r["period"], r[col_name], c(int(r["n"]), font=bk.fin), f"=COUNTIFS({base})",
                       f"=COUNTIFS({base},'By record'!$N:$N,\"critical\")", f"=COUNTIFS({base},'By record'!$N:$N,\"major\")",
                       f"=COUNTIFS({base},'By record'!$N:$N,\"minor\")", c(f"=IF(C{row}>0,D{row}/C{row},0)", fmt="0.0%")])
